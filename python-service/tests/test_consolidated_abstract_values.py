"""Consolidated Abstract mapping from the three branch Financials snapshots."""

from app.engines.financials_engine.engine.consolidated_abstract_values import (
    build_consolidated_abstract_row_measures,
)
from app.engines.financials_engine.engine.consolidated_financials_template import (
    ABSTRACT_SHEET_NAME,
    TSHAPE_SHEET_NAME,
    build_consolidated_financials_workbook_bytes,
)
from io import BytesIO

from openpyxl import load_workbook


def _gold_branch(opening_qty, opening_amt, extra=None):
    payload = {
        'openingPivot': [
            {'product': 'Standard Gold 24K', 'sumOfQuantity': opening_qty, 'sumOfGross': opening_amt}
        ],
        'salesPivot': [],
        'purchasesPivot': [],
        'mrPivots': {},
        'dcPivots': {},
        'layoutByCategory': {},
    }
    if extra:
        payload.update(extra)
    return payload


def _diamond_layout(*, opening_qty, opening_amt, beads_qty=None, beads_amt=None, black_qty=None, black_amt=None):
    rows = [
        {
            'kind': 'grand_total',
            'label': 'GRAND TOTAL',
            'openingQty': opening_qty,
            'openingAmt': opening_amt,
            'purchasesQty': 0,
            'purchasesAmt': 0,
            'salesQty': 10,
            'salesAmt': 100,
            'closingStockQty': opening_qty,
            'closingStockAmt': opening_amt,
        }
    ]
    if beads_qty is not None:
        rows.append(
            {
                'kind': 'subcategory_total',
                'subcategory': 'Diamonds - Beads',
                'openingQty': beads_qty,
                'openingAmt': beads_amt,
            }
        )
    if black_qty is not None:
        rows.append(
            {
                'kind': 'subcategory_total',
                'subcategory': 'Diamonds - Black Diamonds',
                'openingQty': black_qty,
                'openingAmt': black_amt,
            }
        )
    return {'Diamond': rows}


def test_sums_opening_from_all_three_branches():
    by_label, contributing, unresolved = build_consolidated_abstract_row_measures(
        {
            'basheerbagh': _gold_branch(1, 10),
            'kokapet': _gold_branch(2, 20),
            'jubileeHills': _gold_branch(3, 30),
        }
    )
    assert contributing == ['basheerbagh', 'kokapet', 'jubileeHills']
    gold = by_label['Gold Account 24K']['Opening Stock']
    assert gold['qty'] == 6
    assert gold['amt'] == 60
    assert by_label['Platinum Account']['Opening Stock']['qty'] is None
    assert by_label['Gold Account Jadau']['Opening Stock']['qty'] is None
    assert any('Platinum Account' in item for item in unresolved)
    assert any('Jadau' in item for item in unresolved)
    assert any('GP AY 2025-26' in item for item in unresolved)


def test_eliminates_inter_branch_gold_transfers():
    bh = _gold_branch(10, 100)
    bh['dcPivots'] = {
        'jubileeHills': [{'product': 'Standard Gold 24K', 'sumOfQuantity': 5, 'sumOfGross': 0}]
    }
    jh = _gold_branch(4, 40)
    jh['mrPivots'] = {
        'jubileeHills': [{'product': 'Standard Gold 24K', 'sumOfQuantity': 5, 'sumOfGross': 0}]
    }
    by_label, contributing, _unresolved = build_consolidated_abstract_row_measures(
        {'basheerbagh': bh, 'kokapet': _gold_branch(0, 0), 'jubileeHills': jh}
    )
    assert contributing == ['basheerbagh', 'kokapet', 'jubileeHills']
    gold = by_label['Gold Account 24K']
    assert gold['Internal Transfer - Total']['qty'] == 0
    assert gold['Internal Consumption - Total']['qty'] == 0


def test_total_excludes_diamond_detail_rows():
    branches = {
        'basheerbagh': {
            'layoutByCategory': _diamond_layout(
                opening_qty=10, opening_amt=100, beads_qty=4, beads_amt=40
            ),
            'salesPivot': [],
            'purchasesPivot': [],
            'openingPivot': [],
        },
        'kokapet': _gold_branch(0, 0),
        'jubileeHills': _gold_branch(0, 0),
    }
    by_label, _contributing, _unresolved = build_consolidated_abstract_row_measures(branches)
    assert by_label['Diamonds Account']['Opening Stock']['qty'] == 10
    assert by_label['Diamonds - Beads']['Opening Stock']['qty'] == 4
    assert by_label['TOTAL']['Opening Stock']['qty'] == 10
    assert by_label['Diamonds total']['Opening Stock']['qty'] == 4


def test_cy_gp_percent_is_derived_from_consolidated_totals_not_averaged():
    def diamond_sales(opening, sales_qty, sales_amt):
        return {
            'layoutByCategory': {
                'Diamond': [
                    {
                        'kind': 'grand_total',
                        'openingQty': opening,
                        'openingAmt': opening * 10,
                        'purchasesQty': 0,
                        'purchasesAmt': 0,
                        'salesQty': sales_qty,
                        'salesAmt': sales_amt,
                        'closingStockQty': 0,
                        'closingStockAmt': 0,
                    }
                ]
            }
        }

    by_label, _, _ = build_consolidated_abstract_row_measures(
        {
            'basheerbagh': diamond_sales(0, 10, 100),
            'kokapet': diamond_sales(0, 10, 300),
            'jubileeHills': diamond_sales(0, 0, 0),
        }
    )
    gp = by_label['Diamonds Account']['Gross Profit']['amt']
    sales = by_label['Diamonds Account']['Sales - Total']['amt']
    pct = by_label['Diamonds Account']['CY G P %']['amt']
    assert sales == 400
    assert pct == gp / sales * 100


def test_black_diamonds_alias_and_workbook_sheet_names():
    branches = {
        'basheerbagh': _gold_branch(0, 0),
        'kokapet': _gold_branch(0, 0),
        'jubileeHills': {
            'layoutByCategory': _diamond_layout(
                opening_qty=8, opening_amt=80, black_qty=3, black_amt=30
            )
        },
    }
    by_label, _, _ = build_consolidated_abstract_row_measures(branches)
    assert by_label['Diamonds - Black diamonds']['Opening Stock']['qty'] == 3
    raw, contributing, _unresolved = build_consolidated_financials_workbook_bytes(branches)
    assert contributing == ['basheerbagh', 'kokapet', 'jubileeHills']
    workbook = load_workbook(BytesIO(raw))
    assert workbook.sheetnames == [ABSTRACT_SHEET_NAME, TSHAPE_SHEET_NAME]
    ws = workbook[ABSTRACT_SHEET_NAME]
    black_row = None
    for row in range(4, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == 'Diamonds - Black diamonds':
            black_row = row
            break
    assert black_row is not None
    assert ws.cell(row=black_row, column=2).value == 3
    assert ws.cell(row=black_row, column=3).value == 30
    assert workbook[TSHAPE_SHEET_NAME]['A1'].value == 'Particulars'


def test_values_stay_blank_until_all_three_branches_are_present():
    by_label, contributing, unresolved = build_consolidated_abstract_row_measures(
        {
            'basheerbagh': _gold_branch(1, 10),
            'kokapet': _gold_branch(2, 20),
        }
    )
    assert contributing == ['basheerbagh', 'kokapet']
    assert by_label == {}
    assert any('Missing: jubileeHills' in item for item in unresolved)
    raw, _contributing, _unresolved = build_consolidated_financials_workbook_bytes(
        {
            'basheerbagh': _gold_branch(1, 10),
            'kokapet': _gold_branch(2, 20),
        }
    )
    workbook = load_workbook(BytesIO(raw))
    ws = workbook[ABSTRACT_SHEET_NAME]
    gold_row = None
    for row in range(4, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == 'Gold Account 24K':
            gold_row = row
            break
    assert gold_row is not None
    assert ws.cell(row=gold_row, column=2).value is None
    assert ws.cell(row=gold_row, column=3).value is None
