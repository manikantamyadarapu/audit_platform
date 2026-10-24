"""Consolidated Financials workbook: Abstract headers then empty TShape."""

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from app.engines.financials_engine.engine.consolidated_financials_template import (
    ABSTRACT_SHEET_NAME,
    CONSOLIDATED_ABSTRACT_BODY_ROWS,
    CONSOLIDATED_ABSTRACT_MEASURE_COLUMNS,
    CONSOLIDATED_ABSTRACT_TITLE,
    TSHAPE_SHEET_NAME,
    build_consolidated_financials_workbook_bytes,
    consolidated_abstract_column_span,
)


def _workbook_bytes(result):
    return result[0] if isinstance(result, tuple) else result


def _merged_range(ws, row: int, column: int) -> tuple[int, int, int, int] | None:
    for merged in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        if min_row <= row <= max_row and min_col <= column <= max_col:
            return min_col, min_row, max_col, max_row
    return None


def test_consolidated_workbook_has_exactly_abstract_and_tshape():
    raw = _workbook_bytes(build_consolidated_financials_workbook_bytes())
    workbook = load_workbook(BytesIO(raw))
    assert workbook.sheetnames == [ABSTRACT_SHEET_NAME, TSHAPE_SHEET_NAME]
    tshape = workbook[TSHAPE_SHEET_NAME]
    assert tshape['A1'].value == 'Particulars'
    assert tshape['B1'].value == 'Jubilee Hills - Head Office'
    assert tshape['C1'].value == 'Kokapet - Branch'
    assert tshape['D1'].value == 'Basheerbagh - Branch'
    assert tshape['E1'].value == 'Total'
    assert tshape['G1'].value == 'Particulars'


def test_consolidated_abstract_two_level_headers():
    raw = _workbook_bytes(build_consolidated_financials_workbook_bytes())
    ws = load_workbook(BytesIO(raw))[ABSTRACT_SHEET_NAME]
    assert ws.cell(row=1, column=1).value == CONSOLIDATED_ABSTRACT_TITLE
    assert ws.cell(row=2, column=1).value == 'Particulars'
    assert ws.cell(row=3, column=1).value is None
    assert _merged_range(ws, 2, 1) == (1, 2, 1, 3)

    col = 2
    for heading, kind in CONSOLIDATED_ABSTRACT_MEASURE_COLUMNS:
        span = consolidated_abstract_column_span(kind)
        assert ws.cell(row=2, column=col).value == heading
        if kind == 'qty_amt':
            assert ws.cell(row=3, column=col).value == 'Qty in Gms / Cts'
            assert ws.cell(row=3, column=col + 1).value == 'Amount in Rs'
            assert _merged_range(ws, 2, col) == (col, 2, col + 1, 2)
        elif kind == 'amt_only':
            assert ws.cell(row=3, column=col).value == 'Amount in Rs'
            assert _merged_range(ws, 2, col) is None
        else:
            assert ws.cell(row=3, column=col).value is None
            assert _merged_range(ws, 2, col) == (col, 2, col, 3)
        col += span

    assert [heading for heading, _kind in CONSOLIDATED_ABSTRACT_MEASURE_COLUMNS] == [
        'Opening Stock',
        'Purchases - Total',
        'Internal Transfer - Total',
        'Internal Consumption - Total',
        'Making Charges',
        'Sales - Total',
        'Closing Stock - Total',
        'Gross Profit',
        'CY G P %',
        'GP AY 2025-26',
        'Deviation',
    ]


def test_consolidated_abstract_particulars_order_without_values():
    raw = _workbook_bytes(build_consolidated_financials_workbook_bytes())
    ws = load_workbook(BytesIO(raw))[ABSTRACT_SHEET_NAME]

    labels = []
    for offset, (kind, label) in enumerate(CONSOLIDATED_ABSTRACT_BODY_ROWS):
        row = 4 + offset
        assert ws.cell(row=row, column=1).value == label
        if kind != 'spacer':
            labels.append(label)
            assert ws.cell(row=row, column=2).value is None

    assert labels == [
        'Gold Account 24K',
        'Gold Account 22K',
        'Gold Account 18K',
        'Gold Account 14K',
        'Platinum Account',
        'Gold Account Jadau',
        'Silver Account',
        'Diamonds Account',
        'Emeralds',
        'Pearls',
        'Rubies',
        'Color Stones Account',
        'TOTAL',
        'DIAMONDS',
        'Diamonds - Black diamonds',
        'Diamonds - Beads',
        'Diamonds Rosecut diamonds',
        'Diamonds - Flat polki',
        'Uncut - diamonds',
        'Diamonds',
        'Diamonds total',
        'COLOR STONES',
        'Precious Stones',
        'Semi Precious',
        'Synthetic Stones',
        'TOTAL Colour Stones',
    ]
