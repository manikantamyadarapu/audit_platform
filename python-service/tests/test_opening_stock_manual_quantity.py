"""Manual Opening Stock Quantity mapping queue (after automatic matching)."""

from io import BytesIO

from openpyxl import Workbook

from app.engines.financials_engine.engine.opening_stock import validate_opening_stock
from app.engines.financials_engine.engine.opening_stock_manual import (
    attach_manual_quantity_mapping,
)
from app.engines.financials_engine.parsers.opening_stock_loader import (
    load_previous_year_opening_stock,
)


def _product_sheet(ws, qty: float, amount: float = 1.0) -> None:
    ws.append(['Particulars', 'Opening stock', None, 'Closing stock', None])
    ws.append([1, 'Qty', 'Amt.', 'Qty', 'Amt.'])
    ws.append(['Closing Balance', qty, 0, qty, amount])


def _manual_prev_workbook() -> bytes:
    wb = Workbook()
    dia = wb.active
    dia.title = 'Dia'
    dia.append(['Particulars', 'Opening stock', None, 'Purchases', None, 'Closing stock', None])
    dia.append([1, 'Qty', 'Amt.', 'Qty', 'Amt.', 'Qty', 'Amt.'])
    dia.append(['Diamonds - Beads'])
    dia.append(['Di. Beads', 263.03, 100.0, 0, 0, 177.86, 234713.15])

    fp2 = wb.create_sheet('Flat polki FP 2')
    _product_sheet(fp2, 10.0, 50.0)
    fp3 = wb.create_sheet('Flat polki FP 3')
    _product_sheet(fp3, 6.0, 30.0)
    polki = wb.create_sheet('Polki')
    _product_sheet(polki, 5.5, 200.0)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestManualOpeningQuantityQueue:
    def test_unmatched_product_gets_same_subcategory_product_sheets_only(self):
        prev = load_previous_year_opening_stock(_manual_prev_workbook(), 'prev.xlsx')
        result = validate_opening_stock(
            quantity_rows=[{'product': 'Flat polki FP 1', 'openingBalance': 16.0}],
            previous_year_sheets=prev['productIndex'],
            subcategory_products=prev.get('subcategoryProducts'),
            sheet_products=prev.get('sheetProducts'),
        )
        report = attach_manual_quantity_mapping(
            dict(result['report']),
            dedicated_product_sheets=prev['dedicatedProductSheets'],
        )
        queue = report['manualQuantityMappingRequired']
        assert len(queue) == 1
        row = queue[0]
        assert row['product'] == 'Flat polki FP 1'
        assert row['openingQty'] == 16.0
        assert row['subcategory'] == 'Diamonds - Flat polki'
        names = {c['product'] for c in row['candidateProducts']}
        assert names == {'Flat polki FP 2', 'Flat polki FP 3'}
        assert 'Di. Beads' not in names
        assert 'Polki' not in names
        qty_by_name = {c['product']: c['closingQty'] for c in row['candidateProducts']}
        assert qty_by_name['Flat polki FP 2'] == 10.0
        assert qty_by_name['Flat polki FP 3'] == 6.0

    def test_automatically_matched_products_are_excluded(self):
        prev = load_previous_year_opening_stock(_manual_prev_workbook(), 'prev.xlsx')
        result = validate_opening_stock(
            quantity_rows=[{'product': 'Di. Beads', 'openingBalance': 263.03}],
            previous_year_sheets=prev['productIndex'],
            subcategory_products=prev.get('subcategoryProducts'),
            sheet_products=prev.get('sheetProducts'),
        )
        report = attach_manual_quantity_mapping(
            dict(result['report']),
            dedicated_product_sheets=prev['dedicatedProductSheets'],
        )
        products = {r['product'] for r in report['manualQuantityMappingRequired']}
        assert 'Di. Beads' not in products
        assert result['report']['exactMatchedCount'] == 1

    def test_dedicated_sheets_are_collected_separately_from_category_tabs(self):
        prev = load_previous_year_opening_stock(_manual_prev_workbook(), 'prev.xlsx')
        dedicated_names = {row['product'] for row in prev['dedicatedProductSheets']}
        assert 'Flat polki FP 2' in dedicated_names
        assert 'Polki' in dedicated_names
        assert 'Di. Beads' not in dedicated_names
