"""Consolidated TShape mapping from the three branch Trading sources."""

from io import BytesIO

from openpyxl import load_workbook

from app.engines.financials_engine.engine.consolidated_financials_template import (
    ABSTRACT_SHEET_NAME,
    TSHAPE_SHEET_NAME,
    build_consolidated_financials_workbook_bytes,
)
from app.engines.financials_engine.engine.consolidated_tshape_values import (
    TSHAPE_LEFT_ROWS,
    TSHAPE_RIGHT_ROWS,
    build_consolidated_tshape_values,
)


def _diamond_branch(*, opening_amt, purchases_amt=0, sales_amt=0, closing_amt=0, receipts_jh=None, issues_bh=None):
    return {
        'layoutByCategory': {
            'Diamond': [
                {
                    'kind': 'grand_total',
                    'openingAmt': opening_amt,
                    'purchasesAmt': purchases_amt,
                    'salesAmt': sales_amt,
                    'closingStockAmt': closing_amt,
                    'receiptsJubileeHillsAmt': receipts_jh,
                    'issuesBanjaraHillsAmt': issues_bh,
                }
            ]
        }
    }


def test_tshape_branch_columns_use_that_branch_opening_and_sales():
    values, unresolved = build_consolidated_tshape_values(
        {
            'jubileeHills': _diamond_branch(opening_amt=10, sales_amt=40, closing_amt=5),
            'kokapet': _diamond_branch(opening_amt=20, sales_amt=50, closing_amt=8),
            'basheerbagh': _diamond_branch(opening_amt=30, sales_amt=60, closing_amt=9),
        }
    )
    assert values['jubileeHills']['Opening Stock'] == 10
    assert values['kokapet']['Opening Stock'] == 20
    assert values['basheerbagh']['Opening Stock'] == 30
    assert values['total']['Opening Stock'] == 60
    assert values['jubileeHills']['Sales'] == 40
    assert values['total']['Sales'] == 150
    assert values['total']['Closing Stock'] == 22
    assert any('Alloy Purchases' in item for item in unresolved)
    assert values['jubileeHills']['Alloy Purchases'] is None


def test_tshape_total_eliminates_matching_inter_branch_transfers():
    values, _unresolved = build_consolidated_tshape_values(
        {
            'jubileeHills': _diamond_branch(opening_amt=0, issues_bh=15),
            'kokapet': _diamond_branch(opening_amt=0, receipts_jh=15),
            'basheerbagh': _diamond_branch(opening_amt=0),
        }
    )
    assert values['jubileeHills']['Stock Transfer to Branches'] == 15
    assert values['kokapet']['Stock transfer from Head Office'] == 15
    assert values['total']['Stock transfer from Head Office'] == 0
    assert values['total']['Stock Transfer to Branches'] == 0


def test_tshape_workbook_layout_and_sheet_names():
    raw, _contributing, _unresolved = build_consolidated_financials_workbook_bytes(
        {
            'basheerbagh': _diamond_branch(opening_amt=30, sales_amt=90, closing_amt=12),
        }
    )
    workbook = load_workbook(BytesIO(raw))
    assert workbook.sheetnames == [ABSTRACT_SHEET_NAME, TSHAPE_SHEET_NAME]
    ws = workbook[TSHAPE_SHEET_NAME]
    assert [ws.cell(row=1, column=col).value for col in range(1, 6)] == [
        'Particulars',
        'Jubilee Hills - Head Office',
        'Kokapet - Branch',
        'Basheerbagh - Branch',
        'Total',
    ]
    left_labels = [ws.cell(row=row, column=1).value for row in range(2, 9)]
    assert left_labels == list(TSHAPE_LEFT_ROWS)
    assert ws.cell(row=2, column=7).value == TSHAPE_RIGHT_ROWS[0]
    assert ws.cell(row=8, column=7).value == 'Total'
    assert ws.cell(row=2, column=4).value is None
    assert ws.cell(row=2, column=5).value is None
    assert ws.cell(row=2, column=8).value is None
    assert ws.cell(row=2, column=10).value is None
