"""Jubilee Hills blank Financials workbook keeps structure and omits values."""

from io import BytesIO

from openpyxl import load_workbook

from app.engines.financials_engine.engine.closing_stock_template import CLOSING_STOCK_CATEGORIES
from app.engines.financials_engine.engine.jubilee_hills_template import (
    build_jubilee_hills_structure_bytes,
)


def test_jubilee_hills_workbook_is_structure_only():
    raw = build_jubilee_hills_structure_bytes(
        company_name='Eximp',
        address='Jubilee Hills',
        financial_year='AY 2025-26',
    )
    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == [
        *CLOSING_STOCK_CATEGORIES,
        'Trading',
        'Abstract',
    ]

    diamond = wb['Diamond']
    assert diamond['A1'].value == 'Eximp'
    assert diamond['A4'].value == 'DETAILS OF JEWELS CLOSING STOCK - DIAMOND'
    assert diamond['A6'].value == 'Particulars / Product'
    assert diamond.cell(row=7, column=6).value == 'Internal Stock Transfer'
    assert diamond.cell(row=7, column=8).value == 'Basheerbagh'
    assert diamond.cell(row=7, column=10).value == 'Kokapet'
    assert diamond.cell(row=7, column=17).value == 'Internal Stock Transfer'
    assert diamond.cell(row=7, column=19).value == 'Basheerbagh'
    assert diamond.cell(row=7, column=21).value == 'Kokapet'
    labels = [diamond.cell(row=row, column=1).value for row in range(10, diamond.max_row + 1)]
    assert 'Diamonds - Beads' in labels
    assert 'TOTAL' in labels
    assert 'GRAND TOTAL' in labels
    assert 'Di. Beads' not in labels

    precious = wb['Precious and Semi Precious']
    precious_labels = [
        precious.cell(row=row, column=1).value for row in range(10, precious.max_row + 1)
    ]
    assert 'Precious Stones' in precious_labels
    assert 'TOTAL - PRECIOUS STONES' in precious_labels

    emerald = wb['Emerald']
    emerald_labels = [emerald.cell(row=row, column=1).value for row in range(10, emerald.max_row + 1)]
    assert emerald_labels == ['GRAND TOTAL']

    for sheet_name in CLOSING_STOCK_CATEGORIES:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=10, max_row=ws.max_row, min_col=2, max_col=ws.max_column):
            for cell in row:
                assert not isinstance(cell.value, (int, float))
                assert cell.value is None or not str(cell.value).startswith('=')

    trading = wb['Trading']
    trading_labels = [
        trading.cell(row=row, column=1).value for row in range(1, trading.max_row + 1)
    ]
    assert 'GOLD ACCOUNT - 24K' in trading_labels
    assert 'DIAMONDS ACCOUNT' in trading_labels
    assert 'To Opening Stock' in trading_labels
    for row in trading.iter_rows(min_row=1, max_row=trading.max_row, min_col=2, max_col=7):
        for cell in row:
            assert not isinstance(cell.value, (int, float))

    abstract = wb['Abstract']
    assert abstract['A1'].value == 'Trading Account Abstract'
    abstract_labels = [
        abstract.cell(row=row, column=1).value for row in range(1, abstract.max_row + 1)
    ]
    assert 'Gold Account 24K' in abstract_labels
    assert 'Diamonds - Beads' in abstract_labels
    assert 'TOTAL' in abstract_labels
    for row in abstract.iter_rows(min_row=4, max_row=abstract.max_row, min_col=2, max_col=abstract.max_column):
        for cell in row:
            assert not isinstance(cell.value, (int, float))
