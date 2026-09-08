"""MR/DC parse + location pivots (separate from Closing Stock mapping)."""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from app.engines.financials_engine.config.mr_dc_columns import MR_DC_COLUMN_SPEC
from app.engines.financials_engine.engine.mr_dc_pivots import (
    build_location_pivots,
    build_location_pivots_with_report,
    build_mr_dc_pivot_payload,
    classify_mr_dc_location,
)
from app.engines.financials_engine.parsers.mr_dc_loader import load_transfer_workbook
from app.utils.sheet_validation_error import SheetValidationError
from test_closing_stock_product_rule_book import SAMPLE_RULE_BOOK, _product_row

HEADERS = [spec[0] for spec in MR_DC_COLUMN_SPEC.values()]


def _blank_row(**overrides) -> dict:
    row = {name: '' for name in HEADERS}
    row.update(overrides)
    return row


def _workbook_bytes(rows: list[dict], *, columns: list[str] | None = None) -> bytes:
    frame = pd.DataFrame(rows)
    if columns is not None:
        frame = frame[columns]
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        frame.to_excel(writer, index=False)
    return buf.getvalue()


class TestMrDcLoader:
    def test_order_independent_and_case_insensitive_headers(self):
        rows = [
            _blank_row(
                Product='Di. Beads',
                Quantity=2,
                **{'Gross Amount': 100},
                Branch='Jubilee Hills',
            ),
        ]
        shuffled = list(reversed(HEADERS))
        renamed = {name: name.upper() if name != 'UOM' else name for name in shuffled}
        frame_rows = [{renamed[k]: v for k, v in rows[0].items()}]
        payload = _workbook_bytes(frame_rows, columns=list(renamed.values()))
        parsed = load_transfer_workbook(payload, 'mr.xlsx', source_label='MR')
        assert len(parsed) == 1
        assert parsed[0]['product'] == 'Di. Beads'
        assert parsed[0]['quantity'] == 2
        assert parsed[0]['grossAmount'] == 100
        assert parsed[0]['branch'] == 'Jubilee Hills'

    def test_whitespace_in_headers(self):
        padded = [f'  {name}  ' for name in HEADERS]
        row = _blank_row(
            Product='Ring',
            Quantity=1,
            **{'Gross Amount': 50},
            Branch='Kokapet',
        )
        payload = _workbook_bytes([{padded[HEADERS.index(k)]: v for k, v in row.items()}], columns=padded)
        parsed = load_transfer_workbook(payload, 'dc.xlsx', source_label='DC')
        assert parsed[0]['product'] == 'Ring'
        assert parsed[0]['branch'] == 'Kokapet'

    def test_missing_column_raises(self):
        incomplete = [name for name in HEADERS if name != 'Branch']
        payload = _workbook_bytes([_blank_row(Product='X', Quantity=1)], columns=incomplete)
        try:
            load_transfer_workbook(payload, 'mr.xlsx', source_label='MR')
            raise AssertionError('expected missing column error')
        except SheetValidationError as exc:
            assert 'Branch' in str(exc)

    def test_dc_succeeds_without_discount_columns(self):
        columns = [name for name in HEADERS if name not in ('Discount', 'Gross Minus Discount')]
        row = _blank_row(
            Product='DC Ring',
            Quantity=4,
            **{'Gross Amount': 80},
            Branch='Jubilee Hills',
        )
        payload = _workbook_bytes([row], columns=columns)
        parsed = load_transfer_workbook(payload, 'dc.xlsx', source_label='DC')
        assert len(parsed) == 1
        assert parsed[0]['product'] == 'DC Ring'
        assert parsed[0]['quantity'] == 4
        assert parsed[0]['grossAmount'] == 80
        assert parsed[0]['branch'] == 'Jubilee Hills'

    def test_dc_requires_only_product_quantity_gross_branch(self):
        columns = ['Product', 'Quantity', 'Gross Amount', 'Branch']
        payload = _workbook_bytes(
            [
                {
                    'Product': 'Kokapet Item',
                    'Quantity': 2,
                    'Gross Amount': 25,
                    'Branch': 'Kokapet',
                }
            ],
            columns=columns,
        )
        parsed = load_transfer_workbook(payload, 'dc.xlsx', source_label='DC')
        assert parsed[0]['product'] == 'Kokapet Item'
        assert parsed[0]['party'] == ''

    def test_mr_requires_only_product_quantity_gross_branch(self):
        columns = ['Product', 'Quantity', 'Gross Amount', 'Branch']
        payload = _workbook_bytes(
            [
                {
                    'Product': 'MR Ring',
                    'Quantity': 3,
                    'Gross Amount': 40,
                    'Branch': 'Jubilee Hills',
                }
            ],
            columns=columns,
        )
        parsed = load_transfer_workbook(payload, 'mr.xlsx', source_label='MR')
        assert parsed[0]['product'] == 'MR Ring'
        assert parsed[0]['quantity'] == 3
        assert parsed[0]['grossAmount'] == 40
        assert parsed[0]['branch'] == 'Jubilee Hills'
        assert parsed[0]['party'] == ''


class TestMrDcPivots:
    def test_classify_locations(self):
        assert classify_mr_dc_location('JUBILEE HILLS') == 'jubileeHills'
        assert classify_mr_dc_location('Kokapet') == 'kokapet'
        assert classify_mr_dc_location('Internal / Basheerbagh') == 'internalBasheerbagh'
        assert classify_mr_dc_location('', 'Basheerbagh') == 'internalBasheerbagh'
        assert classify_mr_dc_location('Unknown') is None

    def test_mr_and_dc_stay_separate(self):
        mr_rows = [
            {
                'product': 'A',
                'quantity': 1,
                'grossAmount': 10,
                'branch': 'Jubilee Hills',
                'party': '',
            },
            {
                'product': 'A',
                'quantity': 2,
                'grossAmount': 20,
                'branch': 'Jubilee Hills',
                'party': '',
            },
            {
                'product': 'B',
                'quantity': 5,
                'grossAmount': 50,
                'branch': 'Kokapet',
                'party': '',
            },
            {
                'product': 'C',
                'quantity': 3,
                'grossAmount': 30,
                'branch': 'Internal',
                'party': '',
            },
        ]
        dc_rows = [
            {
                'product': 'A',
                'quantity': 9,
                'grossAmount': 90,
                'branch': 'Jubilee Hills',
                'party': '',
            },
        ]
        payload = build_mr_dc_pivot_payload(mr_rows=mr_rows, dc_rows=dc_rows)
        assert payload['mrPivots']['jubileeHills'] == [
            {'product': 'A', 'sumOfQuantity': 3.0, 'sumOfGross': 30.0},
        ]
        assert payload['mrPivots']['kokapet'] == [
            {'product': 'B', 'sumOfQuantity': 5.0, 'sumOfGross': 50.0},
        ]
        assert payload['mrPivots']['internalBasheerbagh'] == [
            {'product': 'C', 'sumOfQuantity': 3.0, 'sumOfGross': 30.0},
        ]
        assert payload['dcPivots']['jubileeHills'] == [
            {'product': 'A', 'sumOfQuantity': 9.0, 'sumOfGross': 90.0},
        ]
        assert payload['dcPivots']['kokapet'] == []
        assert payload['dcPivots']['internalBasheerbagh'] == []

    def test_location_pivots_ignore_unmatched_branch(self):
        rows = [
            {
                'product': 'Z',
                'quantity': 1,
                'grossAmount': 1,
                'branch': 'Other',
                'party': '',
            }
        ]
        assert build_location_pivots(rows) == {
            'jubileeHills': [],
            'kokapet': [],
            'internalBasheerbagh': [],
        }

    def test_classification_report_counts_classified_and_unclassified(self):
        rows = [
            {
                'product': 'A',
                'quantity': 1,
                'grossAmount': 10,
                'branch': 'Jubilee Hills',
                'party': '',
            },
            {
                'product': 'B',
                'quantity': 2,
                'grossAmount': 20,
                'branch': 'Unknown Place',
                'party': '',
            },
            {
                'product': 'C',
                'quantity': 3,
                'grossAmount': 30,
                'branch': 'Kokapet',
                'party': '',
            },
        ]
        pivots, report = build_location_pivots_with_report(rows)
        assert report['sourceRowCount'] == 3
        assert report['classifiedRowCount'] == 2
        assert report['unclassifiedCount'] == 1
        assert report['locationCounts']['jubileeHills'] == 1
        assert report['locationCounts']['kokapet'] == 1
        assert report['unclassifiedRows'][0]['product'] == 'B'
        assert pivots['jubileeHills'][0]['product'] == 'A'

    def test_payload_includes_mr_and_dc_reports(self):
        payload = build_mr_dc_pivot_payload(
            mr_rows=[
                {
                    'product': 'A',
                    'quantity': 1,
                    'grossAmount': 1,
                    'branch': 'Other',
                    'party': '',
                }
            ],
            dc_rows=[
                {
                    'product': 'A',
                    'quantity': 1,
                    'grossAmount': 1,
                    'branch': 'Jubilee Hills',
                    'party': '',
                }
            ],
        )
        assert payload['mrReport']['unclassifiedCount'] == 1
        assert payload['mrReport']['classifiedRowCount'] == 0
        assert payload['dcReport']['classifiedRowCount'] == 1
        assert payload['dcReport']['unclassifiedCount'] == 0
        assert payload['dcPivots']['jubileeHills'][0]['product'] == 'A'


class TestMrDcClosingStockQty:
    """Net quantity only — written onto Closing Stock product rows."""

    def test_net_maps_to_layout_rows_per_branch(self):
        from app.engines.financials_engine.config.product_rule_book import (
            map_pivots_to_closing_stock_categories,
        )
        from app.engines.financials_engine.engine.closing_stock_template import (
            build_closing_stock_template_bytes,
            leaf_index_for_measure,
        )

        mr_pivots = {
            'jubileeHills': [{'product': '  PRODUCT   A ', 'sumOfQuantity': 10, 'sumOfGross': 999}],
            'kokapet': [{'product': 'Product A', 'sumOfQuantity': 2, 'sumOfGross': 1}],
            'internalBasheerbagh': [{'product': 'Product C', 'sumOfQuantity': 9, 'sumOfGross': 1}],
        }
        dc_pivots = {
            'jubileeHills': [{'product': 'product a', 'sumOfQuantity': 3, 'sumOfGross': 1}],
            'kokapet': [{'product': 'Product A', 'sumOfQuantity': 5, 'sumOfGross': 1}],
            'internalBasheerbagh': [{'product': 'Product C', 'sumOfQuantity': 2, 'sumOfGross': 1}],
        }
        original_mr = mr_pivots['jubileeHills'][0].copy()

        result = map_pivots_to_closing_stock_categories(
            sales_pivot=[{'product': 'Product A', 'sumOfQuantity': 1, 'sumOfGross': 10}],
            purchases_pivot=[],
            mr_pivots=mr_pivots,
            dc_pivots=dc_pivots,
            rule_book=SAMPLE_RULE_BOOK,
        )
        product_a = _product_row(result['layoutByCategory']['Diamond'], 'Product A')
        product_c = _product_row(result['layoutByCategory']['Emerald'], 'Product C')

        assert product_a['salesQty'] == 1
        assert product_a['receiptsJubileeHillsQty'] == 7
        assert product_a['issuesBanjaraHillsQty'] is None
        assert product_a['issuesKokapetQty'] == 3
        assert product_a['receiptsKokapetQty'] is None
        assert product_a.get('receiptsJubileeHillsAmt') is None
        assert product_c['receiptsInternalQty'] == 7
        assert product_c['issuesInternalQty'] is None

        product_b = _product_row(result['layoutByCategory']['Diamond'], 'Product B')
        assert product_b['receiptsJubileeHillsQty'] is None
        assert product_b['issuesKokapetQty'] is None

        assert mr_pivots['jubileeHills'][0] == original_mr

        raw = build_closing_stock_template_bytes(
            products_by_category=result['productsByCategory'],
            layout_by_category=result['layoutByCategory'],
        )
        wb = load_workbook(BytesIO(raw))
        diamond = wb['Diamond']
        product_a_row = 11
        assert diamond.cell(row=product_a_row, column=1).value == 'Product A'
        assert diamond.cell(
            row=product_a_row, column=2 + leaf_index_for_measure('receiptsJubileeHillsQty')
        ).value == 7
        assert diamond.cell(
            row=product_a_row, column=2 + leaf_index_for_measure('issuesKokapetQty')
        ).value == 3
        assert diamond.cell(
            row=product_a_row, column=2 + leaf_index_for_measure('receiptsJubileeHillsQty') + 1
        ).value is None
        assert diamond.cell(row=product_a_row, column=2 + 12).value is None

    def test_exact_normalized_match_rejects_fuzzy_and_prefix(self):
        from app.engines.financials_engine.config.product_rule_book import (
            map_pivots_to_closing_stock_categories,
        )

        result = map_pivots_to_closing_stock_categories(
            sales_pivot=[],
            purchases_pivot=[],
            mr_pivots={
                'jubileeHills': [
                    {'product': 'ProductA', 'sumOfQuantity': 4, 'sumOfGross': 0},
                    {'product': 'Product', 'sumOfQuantity': 4, 'sumOfGross': 0},
                    {'product': 'oduct A', 'sumOfQuantity': 4, 'sumOfGross': 0},
                ],
                'kokapet': [],
                'internalBasheerbagh': [],
            },
            dc_pivots={'jubileeHills': [], 'kokapet': [], 'internalBasheerbagh': []},
            rule_book=SAMPLE_RULE_BOOK,
        )
        product_a = _product_row(result['layoutByCategory']['Diamond'], 'Product A')
        assert product_a['receiptsJubileeHillsQty'] is None
        assert product_a['issuesBanjaraHillsQty'] is None

