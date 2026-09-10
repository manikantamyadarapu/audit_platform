"""Final E2E Financials verification — no business-logic changes."""

from __future__ import annotations

import json
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

BASE = "http://127.0.0.1:8000"
OUT = Path(__file__).resolve().parents[1] / "debug" / "e2e_financials_verify"
DOWNLOADS = OUT / "downloads"
DOWNLOADS.mkdir(parents=True, exist_ok=True)


def http_json(method: str, path: str, payload: dict | None = None, timeout: int = 180):
    data = None
    headers = {}
    if payload is not None:
        data = json.dumps(payload).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(f"{BASE}{path}", data=data, headers=headers, method=method)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        body = resp.read()
        ctype = resp.headers.get("Content-Type", "")
        return resp.status, body, ctype


def http_json_obj(method: str, path: str, payload: dict | None = None, timeout: int = 180):
    status, body, ctype = http_json(method, path, payload, timeout=timeout)
    obj = json.loads(body.decode("utf-8")) if "json" in ctype or body[:1] in (b"{", b"[") else {}
    return status, obj


def main() -> None:
    payload = json.loads((OUT / "validate.json").read_text(encoding="utf-8"))
    data = payload.get("data") if isinstance(payload.get("data"), dict) else payload

    print("success:", data.get("success"))
    print("auditKey:", data.get("auditKey"), "fileType:", data.get("fileType"))
    summary = data.get("summary") or {}
    print("SUMMARY:")
    for key in sorted(summary.keys()):
        value = summary[key]
        if isinstance(value, (dict, list)):
            print(f"  {key}: {type(value).__name__} len={len(value)}")
        else:
            print(f"  {key}: {value}")

    sales = data.get("salesPivot") or []
    purch = data.get("purchasesPivot") or []
    opening = data.get("openingPivot") or []
    mr = data.get("mrPivots") or {}
    dc = data.get("dcPivots") or {}
    print(
        "PIVOT COUNTS:",
        {
            "sales": len(sales),
            "purchases": len(purch),
            "opening": len(opening),
            "mr": {k: len(v or []) for k, v in (mr.items() if isinstance(mr, dict) else [])},
            "dc": {k: len(v or []) for k, v in (dc.items() if isinstance(dc, dict) else [])},
        },
    )
    print("mrReport:", data.get("mrReport"))
    print("dcReport:", data.get("dcReport"))

    layouts = data.get("layoutByCategory") or {}
    print("categories:", list(layouts.keys()))
    hits = {
        "opening": 0,
        "purchases": 0,
        "sales": 0,
        "receipts": 0,
        "issues": 0,
        "closing": 0,
    }
    sample = None
    for cat, rows in layouts.items():
        for row in rows or []:
            if row.get("kind") != "product":
                continue
            if row.get("openingQty") is not None:
                hits["opening"] += 1
            if row.get("purchasesQty") is not None:
                hits["purchases"] += 1
            if row.get("salesQty") is not None:
                hits["sales"] += 1
            if any(
                row.get(k) is not None
                for k in (
                    "receiptsJubileeHillsQty",
                    "receiptsKokapetQty",
                    "receiptsInternalQty",
                )
            ):
                hits["receipts"] += 1
            if any(
                row.get(k) is not None
                for k in (
                    "issuesBanjaraHillsQty",
                    "issuesKokapetQty",
                    "issuesInternalQty",
                )
            ):
                hits["issues"] += 1
            if row.get("closingStockQty") is not None:
                hits["closing"] += 1
            if sample is None and row.get("openingQty") is not None and (
                row.get("salesQty") is not None or row.get("receiptsJubileeHillsQty") is not None
            ):
                sample = {
                    "category": cat,
                    "label": row.get("label"),
                    **{
                        k: row.get(k)
                        for k in (
                            "openingQty",
                            "openingAmt",
                            "purchasesQty",
                            "purchasesAmt",
                            "salesQty",
                            "salesAmt",
                            "receiptsJubileeHillsQty",
                            "receiptsKokapetQty",
                            "receiptsInternalQty",
                            "issuesBanjaraHillsQty",
                            "issuesKokapetQty",
                            "issuesInternalQty",
                            "totalQty",
                            "averageRate",
                            "closingStockQty",
                            "closingStockAmt",
                            "grossProfitAmt",
                            "grossProfitPct",
                        )
                    },
                }
    print("products with measures:", hits)
    print("sample mapped product:", sample)

    checked = ok = 0
    mismatches = []
    for _cat, rows in layouts.items():
        for row in rows or []:
            if row.get("kind") != "product" or row.get("closingStockQty") is None:
                continue
            oq = float(row.get("openingQty") or 0) if row.get("openingQty") is not None else None
            pq = float(row.get("purchasesQty") or 0) if row.get("purchasesQty") is not None else None
            receipts = 0
            has_r = False
            for key in (
                "receiptsJubileeHillsQty",
                "receiptsKokapetQty",
                "receiptsInternalQty",
            ):
                if row.get(key) is not None:
                    has_r = True
                    receipts += float(row.get(key) or 0)
            issues_q = 0
            for key in (
                "issuesBanjaraHillsQty",
                "issuesKokapetQty",
                "issuesInternalQty",
            ):
                if row.get(key) is not None:
                    issues_q += float(row.get(key) or 0)
            if oq is None and pq is None and not has_r:
                continue
            checked += 1
            expected = (oq or 0) + (pq or 0) + receipts - issues_q
            actual = float(row["closingStockQty"])
            if abs(expected - actual) < 1e-6:
                ok += 1
            else:
                mismatches.append(
                    {"label": row.get("label"), "expected": expected, "actual": actual}
                )
    print(
        "closingStockQty formula checks:",
        {
            "checked": checked,
            "ok": ok,
            "mismatch_count": len(mismatches),
            "mismatches": mismatches[:5],
        },
    )
    print(
        "preview/layout product rows:",
        sum(1 for rows in layouts.values() for r in (rows or []) if r.get("kind") == "product"),
    )

    def save_xlsx(name: str, content: bytes):
        path = DOWNLOADS / name
        path.write_bytes(content)
        ok_xlsx = content[:2] == b"PK"
        sheets = []
        if ok_xlsx:
            wb = load_workbook(BytesIO(content), read_only=True)
            sheets = wb.sheetnames
            wb.close()
        print(f"DOWNLOAD {name}: bytes={len(content)} xlsx={ok_xlsx} sheets={sheets}")
        return path, sheets

    _status, content, _ctype = http_json(
        "POST",
        "/api/v1/process/financials/export-pivots",
        {"salesPivot": sales, "purchasesPivot": purch},
    )
    save_xlsx("sales_purchases_pivots.xlsx", content)

    _status, content, _ctype = http_json(
        "POST",
        "/api/v1/process/financials/export-pivots",
        {"salesPivot": sales, "purchasesPivot": []},
    )
    save_xlsx("sales_pivot.xlsx", content)

    _status, content, _ctype = http_json(
        "POST",
        "/api/v1/process/financials/export-pivots",
        {"salesPivot": [], "purchasesPivot": purch},
    )
    save_xlsx("purchases_pivot.xlsx", content)

    _status, content, _ctype = http_json(
        "POST",
        "/api/v1/process/financials/export-closing-stock",
        {
            "salesPivot": sales,
            "purchasesPivot": purch,
            "openingPivot": opening,
            "mrPivots": mr,
            "dcPivots": dc,
            "companyName": "Shree Jewellers Eximp",
            "address": "Basheerbagh",
            "financialYear": "AY 2025-26",
        },
    )
    path, _sheets = save_xlsx("stock_reconciliation.xlsx", content)

    wb = load_workbook(BytesIO(path.read_bytes()), data_only=True)
    print("working paper sheets:", wb.sheetnames)
    dia = wb["Diamond"] if "Diamond" in wb.sheetnames else wb[wb.sheetnames[0]]
    filled = 0
    for row in dia.iter_rows(min_row=1, max_row=80, values_only=True):
        if row and row[0] and any(
            isinstance(c, (int, float)) for c in row[1:8] if c is not None
        ):
            filled += 1
    print("Diamond sheet numeric-ish rows in first 80:", filled)
    if sample and sample.get("label"):
        found_label = False
        for row in dia.iter_rows(min_row=1, max_row=200, values_only=True):
            if row and row[0] == sample["label"]:
                found_label = True
                print("workbook row for sample label:", row[:10])
                break
        print("sample label present in Diamond sheet:", found_label)
    wb.close()

    def write_flat(name: str, rows: list):
        wb = Workbook()
        ws = wb.active
        ws.title = "Pivot"
        ws.append(["Product", "Sum of Quantity", "Sum of Gross Amount"])
        for row in rows:
            ws.append([row.get("product"), row.get("sumOfQuantity"), row.get("sumOfGross")])
        path = DOWNLOADS / name
        wb.save(path)
        print(f"DOWNLOAD {name}: rows={len(rows)} bytes={path.stat().st_size}")

    def write_tree(name: str, tree: dict, label: str):
        wb = Workbook()
        first = True
        for loc, title in [
            ("jubileeHills", "Jubilee Hills"),
            ("kokapet", "Kokapet"),
            ("internalBasheerbagh", "Internal Basheerbagh"),
        ]:
            ws = wb.active if first else wb.create_sheet()
            first = False
            ws.title = f"{label} - {title}"[:31]
            ws.append(["Product", "Sum of Quantity", "Sum of Gross Amount"])
            for row in tree.get(loc) or []:
                ws.append([row.get("product"), row.get("sumOfQuantity"), row.get("sumOfGross")])
        path = DOWNLOADS / name
        wb.save(path)
        print(f"DOWNLOAD {name}: bytes={path.stat().st_size} sheets={wb.sheetnames}")

    write_flat("opening_pivot.xlsx", opening)
    write_tree("mr_pivot.xlsx", mr, "MR")
    write_tree("dc_pivot.xlsx", dc, "DC")

    sp = load_workbook(DOWNLOADS / "sales_pivot.xlsx", read_only=True)
    rows = list(sp["Sales Pivot"].iter_rows(values_only=True))
    print("sales_pivot.xlsx header:", rows[0], "data_rows:", max(0, len(rows) - 1))
    sp.close()
    pp = load_workbook(DOWNLOADS / "purchases_pivot.xlsx", read_only=True)
    rows = list(pp["Purchases Pivot"].iter_rows(values_only=True))
    print("purchases_pivot.xlsx header:", rows[0], "data_rows:", max(0, len(rows) - 1))
    pp.close()
    op = load_workbook(DOWNLOADS / "opening_pivot.xlsx", read_only=True)
    rows = list(op.active.iter_rows(values_only=True))
    print("opening_pivot.xlsx data_rows:", max(0, len(rows) - 1))
    op.close()
    mr_wb = load_workbook(DOWNLOADS / "mr_pivot.xlsx", read_only=True)
    print("mr_pivot.xlsx sheets:", mr_wb.sheetnames)
    mr_wb.close()
    dc_wb = load_workbook(DOWNLOADS / "dc_pivot.xlsx", read_only=True)
    print("dc_pivot.xlsx sheets:", dc_wb.sheetnames)
    dc_wb.close()

    status, body = http_json_obj("GET", "/api/v1/process/financials/closing-stock-rule-book")
    print("GET rule-book", status, "keys", list(body.keys())[:8])

    status, body = http_json_obj(
        "POST",
        "/api/v1/process/financials/remap-closing-stock",
        {
            "salesPivot": sales,
            "purchasesPivot": purch,
            "openingPivot": opening,
            "mrPivots": mr,
            "dcPivots": dc,
        },
    )
    print("POST remap-closing-stock", status, "success", body.get("success"))

    _status, oa_bytes, _ctype = http_json("GET", "/openapi.json")
    oa = json.loads(oa_bytes.decode("utf-8"))
    fin = [p for p in oa.get("paths", {}) if "financial" in p.lower()]
    print("OpenAPI financial paths:", len(fin))
    for path in sorted(fin):
        print(" ", path)


if __name__ == "__main__":
    try:
        main()
    except urllib.error.HTTPError as exc:
        print("HTTPError", exc.code, exc.read()[:1000])
        raise
